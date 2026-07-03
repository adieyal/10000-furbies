import csv
from .normalize_month import normalize_month
KEYWORDS=('tourist','visitor','country','nationality','residence','citizenship')
def parse_workbook(path, source_id, default_month=None):
    # Dependency-light parser for CSV-like official exports; XLSX support is enabled when openpyxl is installed.
    rows=[]; path=str(path)
    if path.lower().endswith('.csv'):
        data=list(csv.reader(open(path, newline='')))
        sheets=[('csv',data)]
    else:
        try:
            from openpyxl import load_workbook
        except ModuleNotFoundError:
            return rows
        wb=load_workbook(path, data_only=True, read_only=True)
        sheets=[(ws.title, [[c for c in row] for row in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    for sheet,data in sheets:
        text=' '.join(str(c).lower() for row in data[:12] for c in row if c is not None)
        if not (('tourist' in text or 'visitor' in text) and any(k in text for k in KEYWORDS[2:])): continue
        header_idx=None
        for i,row in enumerate(data[:20]):
            vals=[str(v).lower() for v in row]
            if sum(1 for v in row if v not in (None,'')) > 1 and any(any(k in v for k in KEYWORDS[2:]) for v in vals) and any(any(k in v for k in ['tourist','visitor','arrival','total']) for v in vals): header_idx=i; break
        if header_idx is None: continue
        header=[str(v).strip().lower() for v in data[header_idx]]
        def find(names):
            for i,h in enumerate(header):
                if any(n in h for n in names): return i
        ci=find(['country','nationality','residence','citizenship']); ai=find(['tourist','visitor','arrival','total']); mi=find(['month','date'])
        if ci is None or ai is None: continue
        for row in data[header_idx+1:]:
            if ci>=len(row) or ai>=len(row): continue
            try: val=int(float(str(row[ai]).replace(',','')))
            except Exception: continue
            try: month=normalize_month(row[mi] if mi is not None and mi<len(row) else default_month)
            except Exception: continue
            if row[ci]: rows.append({'month':month,'raw_country':str(row[ci]).strip(),'arrivals':val,'source_id':source_id,'source_file':path,'source_table':sheet,'parser_name':'parse_excel.header_keyword_v1','parse_confidence':0.9,'notes':'table accepted by title/header keywords'})
    return rows
